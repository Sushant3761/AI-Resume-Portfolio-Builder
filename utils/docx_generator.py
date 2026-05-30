import io
import re
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def compile_docx(markdown_text: str) -> bytes:
    """
    Parses Markdown and compiles it into a cleanly formatted, fully editable Word document.
    """
    doc = Document()
    
    # Configure professional 0.75-inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
    lines = markdown_text.split("\n")
    
    for line in lines:
        line = line.strip()
        
        # Heading 1 (Name)
        if line.startswith("# "):
            val = line[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(val)
            run.font.name = "Arial"
            run.font.size = Pt(18)
            run.bold = True
            
        # Heading 2 (Sections)
        elif line.startswith("## "):
            val = line[3:].strip()
            p = doc.add_paragraph()
            # Add line spacing before header
            p.paragraph_format.space_before = Pt(8)
            run = p.add_run(val)
            run.font.name = "Arial"
            run.font.size = Pt(12)
            run.bold = True
            
        # Heading 3 (Role/Company)
        elif line.startswith("### "):
            val = line[4:].strip()
            p = doc.add_paragraph()
            run = p.add_run(val)
            run.font.name = "Arial"
            run.font.size = Pt(10.5)
            run.bold = True
            
        # Bullet points
        elif line.startswith("* ") or line.startswith("- "):
            val = line[2:].strip()
            clean_val = re.sub(r'\*\*(.*?)\*\*', r'\1', val)
            doc.add_paragraph(clean_val, style='List Bullet')
            
        # Empty space
        elif not line:
            pass
            
        # Normal body paragraphs
        else:
            clean_line = re.sub(r'\*\*(.*?)\*\*', r'\1', line)
            p = doc.add_paragraph(clean_line)
            p.paragraph_format.line_spacing = 1.15
            for run in p.runs:
                run.font.name = "Arial"
                run.font.size = Pt(10)
                
    # Save to a memory stream
    file_stream = io.BytesIO()
    doc.save(file_stream)
    return file_stream.getvalue()
